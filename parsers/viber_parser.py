"""Viber signature carving with bounded blocks and streaming XLSX output."""
import argparse
import math
import os
import re
import tempfile
from pathlib import Path

try:
    from .parser_io import mapped_input, check_output
except ImportError:
    from parser_io import mapped_input, check_output

PATTERN1 = bytes.fromhex("1A000902005708000008000808080808813908090800000800000081")
PATTERN2 = bytes.fromhex("7B226465736B746F705F696E666F223A7B22686173466F7277617264496E6469636174696F6E223A66616C73652C22696E69746961746F72223A747275657D2C227574634F666673657453656373223A33323430307D")
PATTERN3 = bytes.fromhex("1A0009020057080000080F0808080808813908090800000800000081")
START_RE = re.compile(b"(?:" + re.escape(PATTERN1) + b"|" + re.escape(PATTERN3) + b")")
ILLEGAL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")

def iter_viber_messages(data, max_block_bytes=65536, max_text_chars=None):
    if max_block_bytes <= 0 or (max_text_chars is not None and max_text_chars < 0):
        raise ValueError("Invalid block/text length limit")
    matches = START_RE.finditer(data)
    current = next(matches, None)
    while current is not None:
        following = next(matches, None)
        text_start = current.end()
        limit = min(len(data), text_start + max_block_bytes + len(PATTERN2))
        if following is not None:
            limit = min(limit, following.start())
        end = data.find(PATTERN2, text_start, limit)
        if end != -1:
            raw = data[text_start:end]
            try:
                text = raw.decode("utf-8")
                decode_errors = False
            except UnicodeDecodeError:
                text = raw.decode("utf-8", errors="replace")
                decode_errors = True
            if text and (max_text_chars is None or len(text) <= max_text_chars):
                yield {
                    'offset': f"0x{current.start():08X}", 'size': len(raw),
                    'hex': raw.hex(), 'text': text,
                    'text_offset': f"0x{text_start:08X}", 'decode_errors': decode_errors,
                }
        current = following

def parse_viber_messages(data, max_block_bytes=65536, max_text_chars=None):
    """Compatibility list API. CLI uses the streaming iterator."""
    return list(iter_viber_messages(data, max_block_bytes, max_text_chars))

def save_to_excel(messages, output_path):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    destination = Path(output_path).resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook(write_only=True)
    headers = ['No', 'Offset', 'Size (bytes)', 'Hex Data', 'Text Data',
               'Text Offset', 'Decode Errors', 'Part', 'Parts']
    ws = wb.create_sheet("Viber Messages")
    ws.append(headers)
    row_count, count, sheet_no = 1, 0, 1
    # Split long cells into explicit continuation rows; hex always preserves bytes.
    for count, msg in enumerate(messages, 1):
        text = ILLEGAL_RE.sub("", msg['text'])
        raw_hex = msg['hex']
        parts = max(1, math.ceil(len(text) / 30000), math.ceil(len(raw_hex) / 30000))
        for part in range(parts):
            if row_count >= 1048576:
                sheet_no += 1
                ws = wb.create_sheet(f"Viber Messages {sheet_no}")
                ws.append(headers)
                row_count = 1
            values = [count, msg['offset'], msg['size'],
                      raw_hex[part*30000:(part+1)*30000], text[part*30000:(part+1)*30000],
                      msg.get('text_offset', ''), msg.get('decode_errors', False),
                      part + 1, parts]
            cells = []
            for value in values:
                cell = WriteOnlyCell(ws, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cells.append(cell)
            ws.append(cells)
            row_count += 1
    fd, temporary = tempfile.mkstemp(prefix=".viber-", suffix=".xlsx", dir=destination.parent)
    os.close(fd)
    try:
        wb.save(temporary)
        os.replace(temporary, destination)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
        wb.close()
    return count

def process_folder(input_folder, output_folder, max_block_bytes=65536, max_text_chars=None):
    source, destination = Path(input_folder).resolve(), Path(output_folder).resolve()
    if source == destination:
        raise ValueError("Input and output folders must be different")
    if not source.is_dir():
        raise ValueError(f"Input folder not found: {source}")
    # Freeze the input list before creating outputs; retain extensions on stem collisions.
    files = sorted(p for p in source.iterdir() if p.is_file())
    stems = {}
    for path in files:
        stems[path.stem.casefold()] = stems.get(path.stem.casefold(), 0) + 1
    failures = 0
    for path in files:
        name = path.name if stems[path.stem.casefold()] > 1 else path.stem
        out = destination / (name + '.xlsx')
        try:
            check_output(path, out)
            with mapped_input(path) as data:
                count = save_to_excel(iter_viber_messages(data, max_block_bytes, max_text_chars), out)
            print(f"{path.name}: {count} records -> {out}")
        except (OSError, ValueError) as exc:
            failures += 1
            print(f"{path.name}: ERROR: {exc}")
    return failures

def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("input_folder")
    ap.add_argument("output_folder")
    ap.add_argument("--max-block-bytes", type=int, default=65536)
    ap.add_argument("--max-text-chars", type=int, default=None,
                    help="Optional legacy filter; use 40 to reproduce the old length cap")
    args = ap.parse_args()
    try:
        failures = process_folder(args.input_folder, args.output_folder,
                                  args.max_block_bytes, args.max_text_chars)
    except (OSError, ValueError) as exc:
        ap.exit(1, f"Error: {exc}\n")
    if failures:
        ap.exit(1, f"{failures} file(s) failed\n")

if __name__ == "__main__":
    main()
